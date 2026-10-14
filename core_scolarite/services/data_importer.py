"""
Service d'import massif (ETL) — core_scolarite (ARCHITECTURE §8.1).

Contrat :
  * 6 en-têtes obligatoires sur la ligne 1 : MATRICULE, NOM, PRENOM, EMAIL,
    FACULTE, DEPARTEMENT (comparaison insensible à la casse et aux espaces).
  * Validation ligne à ligne ; toutes les erreurs sont collectées et renvoyées
    d'un bloc (`ImportValidationError.errors` = [{line, message}]).
  * Une seule erreur ⇒ **rollback total** (`transaction.atomic`) : aucun
    enregistrement partiel, jamais.
  * `strict_placement=False` (mode MVP) tolère faculté/département vides →
    profil créé non onboardé (Cas A). En mode `True`, les deux sont exigés et
    le couple (faculté, département) doit être cohérent.

Le fichier accepte les extensions `.csv` (utf-8-sig) et `.xlsx` (openpyxl).
"""

from __future__ import annotations

from datetime import date
from typing import Literal

from django.conf import settings
from django.contrib.auth.models import User
from django.core.exceptions import ValidationError as DjangoValidationError
from django.core.validators import validate_email
from django.db import transaction
from openpyxl import load_workbook

from accounts.models import Department, Faculty, Student, Teacher

REQUIRED_HEADERS = ["MATRICULE", "NOM", "PRENOM", "EMAIL", "FACULTE", "DEPARTEMENT"]

Entity = Literal["students", "teachers"]


class ImportValidationError(Exception):
    """Erreur d'import bloquante (en-tête manquant ou erreurs de lignes)."""

    def __init__(self, message: str, errors: list[dict] | None = None):
        self.message = message
        self.errors = errors or []
        super().__init__(message)


def _normalize_headers(raw_headers) -> list[str]:
    return [(str(h) if h is not None else "").strip().upper() for h in raw_headers]


def _check_headers(headers: list[str]) -> None:
    missing = [h for h in REQUIRED_HEADERS if h not in headers]
    if missing:
        raise ImportValidationError(
            "Colonne(s) manquante(s) dans l'en-tête : "
            + ", ".join(missing)
            + f". En-têtes attendus : {', '.join(REQUIRED_HEADERS)}."
        )


def _tabular_rows(rows_source):
    """Consomme les en-têtes puis renvoie (headers, lignes) normalisés."""
    iterator = iter(rows_source)
    try:
        raw_headers = next(iterator)
    except StopIteration:
        raise ImportValidationError("Le fichier est vide.")
    headers = _normalize_headers(raw_headers)

    data: list[tuple[int, dict]] = []
    for line_number, values in enumerate(iterator, start=2):
        if values is None:
            continue
        values = list(values)
        if all(v is None or str(v).strip() == "" for v in values):
            continue  # ignore les lignes totalement vides
        record = {}
        for index, header in enumerate(headers):
            value = values[index] if index < len(values) else None
            record[header] = str(value).strip() if value is not None else ""
        data.append((line_number, record))
    return headers, data


def _load_rows(source, name: str):
    """Renvoie (headers, lignes) selon l'extension du fichier."""
    lower = name.lower()
    if lower.endswith(".csv"):
        import csv
        import io

        text = source.read()
        if isinstance(text, bytes):
            text = text.decode("utf-8-sig")
        return _tabular_rows(csv.reader(io.StringIO(text)))
    if lower.endswith(".xlsx"):
        workbook = load_workbook(source, read_only=True, data_only=True)
        try:
            sheet = workbook.active
            return _tabular_rows(sheet.iter_rows(values_only=True))
        finally:
            workbook.close()
    raise ImportValidationError(
        "Format de fichier non supporté : utilisez un fichier .csv ou .xlsx."
    )


def _is_valid_email(value: str) -> bool:
    try:
        validate_email(value)
        return True
    except DjangoValidationError:
        return False


@transaction.atomic
def import_users(source, entity: Entity = "students", strict_placement: bool = True) -> int:
    """
    Importe des étudiants ou enseignants depuis un flux CSV/XLSX.

    :param source: objet fichier (UploadedFile ou similaire).
    :param entity: "students" ou "teachers".
    :param strict_placement: exiger faculté + département (True) ou les tolérer (False).
    :return: nombre d'enregistrements créés.
    :raises ImportValidationError: en-têtes manquants ou lignes invalides (rollback total).
    """
    if entity not in ("students", "teachers"):
        raise ImportValidationError("Entité inconnue : utilisez 'students' ou 'teachers'.")

    if entity == "teachers":
        strict_placement = True  # les enseignants exigent un rattachement complet

    name = getattr(source, "name", "") or ""
    headers, rows = _load_rows(source, name)
    _check_headers(headers)
    if not rows:
        raise ImportValidationError("Le fichier ne contient aucune ligne de données.")

    # Pré-fetch (2 requêtes, pas de N+1) : facultés et couples (faculté, département).
    faculties = {f.name.strip().upper(): f for f in Faculty.objects.all()}
    departments = {(d.faculty_id, d.name.strip().upper()): d for d in Department.objects.all()}
    existing_matricules = set(
        User.objects.filter(username__in=[r["MATRICULE"] for _, r in rows]).values_list(
            "username", flat=True
        )
    )

    errors: list[dict] = []
    seen_matricules: set[str] = set()
    valid_rows: list[tuple[int, dict, Faculty | None, Department | None]] = []

    for line_number, record in rows:
        matricule = (record.get("MATRICULE") or "").strip()
        nom = (record.get("NOM") or "").strip()
        prenom = (record.get("PRENOM") or "").strip()
        email = (record.get("EMAIL") or "").strip()
        faculty_name = (record.get("FACULTE") or "").strip()
        department_name = (record.get("DEPARTEMENT") or "").strip()

        line_errors = []
        if not matricule:
            line_errors.append("MATRICULE vide.")
        if not nom:
            line_errors.append("NOM vide.")
        if not prenom:
            line_errors.append("PRENOM vide.")
        if not email:
            line_errors.append("EMAIL vide.")
        elif not _is_valid_email(email):
            line_errors.append(f"EMAIL invalide : {email}.")
        if matricule:
            if matricule in existing_matricules:
                line_errors.append(f"MATRICULE déjà existant : {matricule}.")
            if matricule in seen_matricules:
                line_errors.append(f"MATRICULE en doublon dans le fichier : {matricule}.")

        # Résolution hiérarchique faculté → département.
        faculty = None
        department = None
        if faculty_name:
            faculty = faculties.get(faculty_name.upper())
            if faculty is None:
                line_errors.append(f"FACULTE inexistante : {faculty_name}.")
        elif strict_placement:
            line_errors.append("FACULTE vide.")

        if department_name:
            if faculty is None:
                line_errors.append(f"DEPARTEMENT « {department_name} » sans FACULTE valide.")
            else:
                department = departments.get((faculty.pk, department_name.upper()))
                if department is None:
                    line_errors.append(
                        f"DEPARTEMENT « {department_name} » non rattaché à « {faculty.name} »."
                    )
        elif strict_placement:
            line_errors.append("DEPARTEMENT vide.")

        if line_errors:
            errors.append({"line": line_number, "message": " ".join(line_errors)})
        else:
            if matricule:
                seen_matricules.add(matricule)
            valid_rows.append((line_number, record, faculty, department))

    if errors:
        raise ImportValidationError(
            f"{len(errors)} ligne(s) invalide(s). Aucun enregistrement n'a été importé.",
            errors=errors,
        )

    # ── Création (tout est déjà validé ; le bloc est atomique) ──────────
    year = date.today().year
    pattern = settings.IMPORT_DEFAULT_PASSWORD_PATTERN
    created = 0

    for _line_number, record, faculty, department in valid_rows:
        matricule = record["MATRICULE"]
        user = User.objects.create(
            username=matricule,
            first_name=record["PRENOM"],
            last_name=record["NOM"],
            email=record["EMAIL"],
        )
        user.set_password(pattern.format(matricule=matricule, year=year))
        user.save(update_fields=["password"])

        if entity == "students":
            Student.objects.create(
                user=user,
                registration_number=matricule,
                faculty=faculty,
                department=department,
            )
        else:
            Teacher.objects.create(user=user, faculty=faculty, department=department)
        created += 1

    return created
